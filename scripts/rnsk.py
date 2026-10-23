#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""rnsk-tutor 的文件与抓取工具。

脚本只处理网络抓取、目录、题目册和错题集；讲课、出题和判分由助手完成。

用法：
    rnsk.py init        --dir <工作目录> --lang python3 [--url <教程入口>] [--limit N] [--force]
    rnsk.py next        --dir <工作目录>
    rnsk.py mark        --dir <工作目录> --board "<板块名字>" [--undo]
    rnsk.py addq        --dir <工作目录> --board <板块> --type 单选题 --question <题干> \
                        --options "A.xx|B.xx|C.xx|D.xx" --answer A
    rnsk.py review      --dir <工作目录> [--board <板块>] [--threshold 3] [--all]
    rnsk.py review-pass --dir <工作目录> --question <题干> [--board <板块>] [--threshold 3]
    rnsk.py review-fail --dir <工作目录> --question <题干> [--board <板块>] [--threshold 3]
    rnsk.py stats       --dir <工作目录>

错题集里的「通过次数」记录连续复习通过次数：复习答对 +1，答错清零；
达到阈值（默认 3）后就不再进入复习列表，直到再次答错被清零。
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
HEADERS = {"User-Agent": UA, "Accept-Language": "zh-CN,zh;q=0.9"}

INDEX_NAME = "知识目录.md"
QUESTION_BOOK = "题目册.xlsx"
WRONG_BOOK = "错题集.xlsx"
Q_SHEET = "题目册"
W_SHEET = "错题集"

Q_HEADERS = ["序号", "板块", "知识点", "难度", "题型", "题目", "选项", "正确答案", "出题时间"]
W_HEADERS = ["回答错误时间", "板块", "题目", "选项", "正确答案", "通过次数"]
DEFAULT_REVIEW_THRESHOLD = 3

# 菜鸟教程常用语言入口
ENTRY = {
    "python": "https://www.runoob.com/python3/python3-tutorial.html",
    "python3": "https://www.runoob.com/python3/python3-tutorial.html",
    "java": "https://www.runoob.com/java/java-tutorial.html",
    "c": "https://www.runoob.com/cprogramming/c-tutorial.html",
    "cpp": "https://www.runoob.com/cplusplus/cpp-tutorial.html",
    "c++": "https://www.runoob.com/cplusplus/cpp-tutorial.html",
    "csharp": "https://www.runoob.com/csharp/csharp-tutorial.html",
    "c#": "https://www.runoob.com/csharp/csharp-tutorial.html",
    "javascript": "https://www.runoob.com/js/js-tutorial.html",
    "js": "https://www.runoob.com/js/js-tutorial.html",
    "typescript": "https://www.runoob.com/typescript/ts-tutorial.html",
    "go": "https://www.runoob.com/go/go-tutorial.html",
    "golang": "https://www.runoob.com/go/go-tutorial.html",
    "rust": "https://www.runoob.com/rust/rust-tutorial.html",
    "php": "https://www.runoob.com/php/php-tutorial.html",
    "sql": "https://www.runoob.com/sql/sql-tutorial.html",
    "mysql": "https://www.runoob.com/mysql/mysql-tutorial.html",
    "html": "https://www.runoob.com/html/html-tutorial.html",
    "css": "https://www.runoob.com/css/css-tutorial.html",
    "kotlin": "https://www.runoob.com/kotlin/kotlin-tutorial.html",
    "swift": "https://www.runoob.com/swift/swift-tutorial.html",
    "ruby": "https://www.runoob.com/ruby/ruby-tutorial.html",
    "lua": "https://www.runoob.com/lua/lua-tutorial.html",
    "perl": "https://www.runoob.com/perl/perl-tutorial.html",
    "django": "https://www.runoob.com/django/django-tutorial.html",
    "regex": "https://www.runoob.com/regexp/regexp-tutorial.html",
    "linux": "https://www.runoob.com/linux/linux-tutorial.html",
    "shell": "https://www.runoob.com/linux/linux-shell.html",
    "git": "https://www.runoob.com/git/git-tutorial.html",
    "docker": "https://www.runoob.com/docker/docker-tutorial.html",
    "redis": "https://www.runoob.com/redis/redis-tutorial.html",
    "pandas": "https://www.runoob.com/pandas/pandas-tutorial.html",
    "numpy": "https://www.runoob.com/numpy/numpy-tutorial.html",
    "mongodb": "https://www.runoob.com/mongodb/mongodb-tutorial.html",
    "xml": "https://www.runoob.com/xml/xml-tutorial.html",
}

STOPWORDS = {"教程", "参考手册", "菜鸟教程", "简介", "实例", "说明", "概述"}

PRETTY = {
    "python": "Python", "python3": "Python3", "java": "Java", "c": "C",
    "cpp": "C++", "c++": "C++", "csharp": "C#", "c#": "C#",
    "javascript": "JavaScript", "js": "JavaScript", "typescript": "TypeScript",
    "go": "Go", "golang": "Go", "rust": "Rust", "php": "PHP", "sql": "SQL",
    "mysql": "MySQL", "html": "HTML", "css": "CSS", "kotlin": "Kotlin",
    "swift": "Swift", "ruby": "Ruby", "lua": "Lua", "perl": "Perl",
    "linux": "Linux", "shell": "Shell", "git": "Git", "docker": "Docker",
    "redis": "Redis", "pandas": "Pandas", "numpy": "NumPy",
    "mongodb": "MongoDB", "xml": "XML", "django": "Django", "regex": "正则表达式",
}


def pretty_lang(lang: str) -> str:
    key = (lang or "").strip().lower()
    return PRETTY.get(key, lang or "编程语言")


# ---------------------------------------------------------------- 基础工具

def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def rnsk_root(work_dir: str) -> str:
    return os.path.join(os.path.abspath(os.path.expanduser(work_dir)), "rnsk")


def fetch(url: str, timeout: int = 25, retries: int = 3) -> str:
    last = None
    for i in range(retries):
        try:
            r = requests.get(url, headers=HEADERS, timeout=timeout)
            r.raise_for_status()
            r.encoding = r.apparent_encoding or "utf-8"
            return r.text
        except Exception as exc:  # noqa: BLE001
            last = exc
            time.sleep(1 + i)
    raise RuntimeError(f"抓取失败 {url}: {last}")


def entry_url(lang: str) -> str:
    key = (lang or "").strip().lower()
    if key in ENTRY:
        return ENTRY[key]
    slug = re.sub(r"[^a-z0-9]", "", key)
    if not slug:
        raise SystemExit("请用 --lang 指定语言，或直接用 --url 给出菜鸟教程入口地址")
    return f"https://www.runoob.com/{slug}/{slug}-tutorial.html"


# ---------------------------------------------------------------- 解析

def parse_toc(entry: str) -> list[dict]:
    html = fetch(entry)
    soup = BeautifulSoup(html, "html.parser")
    box = soup.find(id="leftcolumn") or soup.find(class_="design") or soup
    base_dir = entry.rsplit("/", 1)[0] + "/"
    items, seen = [], set()
    for a in box.find_all("a", href=True):
        title = a.get_text(" ", strip=True)
        if not title:
            continue
        url = urljoin(entry, a["href"]).split("#")[0]
        if not url.startswith(base_dir) or not url.endswith(".html"):
            continue
        if url in seen:
            continue
        seen.add(url)
        items.append({"title": title, "url": url})
    return items


def clean_desc(title: str, desc: str) -> str:
    """把 meta description 收拾成一句不超过 60 字的板块概述。"""
    text = re.sub(r"\s+", " ", desc or "").strip()
    text = re.sub(r"^[A-Za-z0-9#+./]{1,12}\s+", "", text)
    pieces = [p for p in re.split(r"\s+", title or "") if len(p) >= 2]
    pieces.append((title or "").replace(" ", ""))
    candidate = text
    changed = True
    while changed and candidate:
        changed = False
        for p in pieces:
            if p and candidate.lower().startswith(p.lower()):
                candidate = candidate[len(p):].strip(" :：，,。")
                changed = True
    candidate = re.sub(r"^[A-Za-z0-9#+.]{1,10}\s*[:：]\s*", "", candidate)
    candidate = re.sub(r"^[的与和了，,。:：]+", "", candidate)
    # 洗过头（开头残缺）就别洗了，退回原文
    if len(candidate) < 12 or candidate[:1] in "的了，。":
        candidate = text
    candidate = re.sub(r"[\s:：]+$", "", candidate)
    if len(candidate) < 8:
        return f"{title} 的核心知识点讲解"
    return candidate[:60] + ("…" if len(candidate) > 60 else "")


def page_info(url: str, title: str = "") -> dict:
    html = fetch(url)
    soup = BeautifulSoup(html, "html.parser")
    h1 = soup.find("h1")
    page_title = title or (h1.get_text(" ", strip=True) if h1 else "")
    body = (soup.find("div", class_="article-body")
            or soup.find(id="content") or soup)
    heads = [h.get_text(" ", strip=True) for h in body.find_all(["h2", "h3"])]
    heads = [h for h in heads if h][:8]
    meta = soup.find("meta", attrs={"name": "description"})
    desc = (meta.get("content") or "").strip() if meta else ""
    if not desc:
        p = body.find("p")
        desc = p.get_text(" ", strip=True) if p else ""
    return {"heads": heads, "desc": clean_desc(page_title, desc)}


def extract_text(url: str, limit: int = 6000) -> str:
    html = fetch(url)
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header",
                     "iframe", "noscript", "form"]):
        tag.decompose()
    body = (soup.find("div", class_="article-body")
            or soup.find(id="content") or soup.body)
    if body is None:
        return ""
    for junk in body.select(".ad, .advertisement, .comment, #comment, .sidebar"):
        junk.decompose()
    text = body.get_text("\n", strip=True)
    text = re.sub(r"\n{2,}", "\n", text)
    return text[:limit]


def gen_points(title: str, heads: list[str]) -> str:
    """从页面小标题里拧出几个干净的中文关键词当要点。"""
    words: list[str] = []
    for raw in heads:
        w = re.sub(r"^[A-Za-z0-9#+./]{1,12}\s*", "", raw).strip()
        w = re.sub(r"^[的与和是]+", "", w)
        w = re.sub(r"[\s:：，,。.]+$", "", w)
        if not (2 <= len(w) <= 8):
            continue
        if not re.search(r"[\u4e00-\u9fff]", w):
            continue
        if any(ch in w for ch in "&:：/|()（）[]"):
            continue
        if w in STOPWORDS or w in words:
            continue
        words.append(w)
        if len(words) >= 4:
            break
    if not words:
        words = [re.sub(r"^[A-Za-z0-9#+./]{1,12}\s*", "", title).strip() or title]
    return " ".join(words)[:24]


# ---------------------------------------------------------------- 知识目录

PROMO_WORDS = ("购买", "付费", "课程", "会员", "广告", "推广", "优惠", "立即报名", "讲师")
ENV_WORDS = ("环境", "安装", "配置", "下载", "编辑器", "编译器", "IDE", "VScode", "VSCode")


def classify_board(title: str, summary: str = "") -> tuple[str, str]:
    text = f"{title} {summary}"
    if any(word.lower() in text.lower() for word in PROMO_WORDS):
        return "推广/导读", "true"
    if any(word.lower() in text.lower() for word in ENV_WORDS):
        return "环境配置", "true"
    return "核心知识", "false"


def render_index(lang: str, boards: list[dict]) -> str:
    lines = [
        f"# {pretty_lang(lang)} 知识目录",
        "",
        f"> 来源：菜鸟教程 runoob.com · 生成时间 {now_str()}",
        "> 学完一个板块就把 [是否完成] 改成 true。",
        "",
    ]
    for b in boards:
        lines += [
            f"[板块名字]: {b['title']}",
            "[是否完成]: false",
            f"[板块类别]: {b.get('category', '核心知识')}",
            f"[建议跳过]: {b.get('skip', 'false')}",
            f"[板块地址]: {b['url']}",
            f"[板块要点]: {b.get('points', '')}",
            f"[板块内容]: {b.get('summary', '')}",
            "",
        ]
    return "\n".join(lines)


def parse_index(path: str) -> list[dict]:
    boards: list[dict] = []
    cur: dict = {}
    with open(path, encoding="utf-8") as fh:
        for raw in fh:
            line = raw.strip()
            if not line:
                if cur:
                    boards.append(cur)
                    cur = {}
                continue
            m = re.match(r"^\[(.+?)\]:\s*(.*)$", line)
            if m:
                cur[m.group(1)] = m.group(2)
    if cur:
        boards.append(cur)
    return boards


def is_done(board: dict) -> bool:
    return str(board.get("是否完成", "false")).strip().lower() == "true"


# ---------------------------------------------------------------- Excel

def _migrate_book(path: str, sheet: str, headers: list[str],
                  defaults: dict | None = None) -> None:
    """给旧表格补充缺失字段，不改已有内容（新列补默认值）。"""
    from openpyxl import load_workbook
    defaults = defaults or {}
    wb = load_workbook(path)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.active
    existing = [c.value for c in ws[1]]
    missing = [h for h in headers if h not in existing]
    if not missing:
        return
    for h in missing:
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=h)
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=col).value is None:
                ws.cell(row=r, column=col, value=defaults.get(h, ""))
    wb.save(path)


def ensure_books(root: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    for fname, sheet, headers, defaults in (
        (QUESTION_BOOK, Q_SHEET, Q_HEADERS, {}),
        (WRONG_BOOK, W_SHEET, W_HEADERS, {"通过次数": 0}),
    ):
        path = os.path.join(root, fname)
        if os.path.exists(path):
            _migrate_book(path, sheet, headers, defaults)
            continue
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        widths = {"序号": 8, "板块": 22, "题型": 10, "题目": 52, "选项": 48,
                  "正确答案": 14, "出题时间": 20, "回答错误时间": 20,
                  "通过次数": 10}
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(h, 18)
        wb.save(path)


def open_book(path: str, sheet: str):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.active
    return wb, ws


def cmd_addq(args) -> None:
    root = rnsk_root(args.dir)
    ensure_books(root)
    path = os.path.join(root, QUESTION_BOOK)
    wb, ws = open_book(path, Q_SHEET)
    headers = [c.value for c in ws[1]]
    missing = [h for h in Q_HEADERS if h not in headers]
    if missing:
        raise ValueError(f"题库缺少字段: {', '.join(missing)}")
    if args.type not in {"单选题", "多选题", "填空题"}:
        raise ValueError("题型只能是单选题、多选题或填空题")
    question = args.question.strip()
    if not question or len(question) > 160:
        raise ValueError("题目不能为空且最多 160 字")
    options = [x.strip() for x in (args.options or "").split("|") if x.strip()]
    answer = args.answer.strip().upper()
    if args.type in {"单选题", "多选题"}:
        if len(options) != 4 or any(not re.match(r"^[A-D]\.", x) for x in options):
            raise ValueError("选择题必须提供 4 个 A-D 选项，格式为 A.xxx|B.xxx|C.xxx|D.xxx")
        letters = set(re.findall(r"[A-D]", answer))
        if not letters or any(ch not in "ABCD" for ch in answer) or len(letters) != len(answer):
            raise ValueError("答案只能是不重复的 A-D 字母")
        if args.type == "单选题" and len(answer) != 1:
            raise ValueError("单选题答案只能有一个字母")
    elif options or not answer:
        raise ValueError("填空题不要选项，且必须填写答案")
    if args.difficulty not in {"基础", "进阶", "综合"}:
        raise ValueError("难度只能是基础、进阶或综合")
    columns = {name: i for i, name in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if (row[columns["板块"]] == args.board
                and row[columns["题目"]] == question):
            raise ValueError("同一板块已有相同题目，拒绝重复写入")
    seq = ws.max_row
    record = {
        "序号": seq, "板块": args.board, "知识点": args.point.strip(),
        "难度": args.difficulty, "题型": args.type, "题目": question,
        "选项": "|".join(options), "正确答案": answer, "出题时间": now_str(),
    }
    ws.append([record.get(h, "") for h in headers])
    wb.save(path)
    print(json.dumps({"ok": True, "序号": seq, "文件": path}, ensure_ascii=False))


def _wrong_rows(ws) -> tuple[dict, list[dict]]:
    """读出错题集，返回（列名->序号, 行列表）。"""
    headers = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(headers)}
    rows = []
    for r in range(2, ws.max_row + 1):
        data = {h: ws.cell(row=r, column=col[h] + 1).value for h in headers}
        rows.append({"row": r, "data": data})
    return col, rows


def _as_int(value, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def cmd_review(args) -> None:
    root = rnsk_root(args.dir)
    path = os.path.join(root, WRONG_BOOK)
    if not os.path.exists(path):
        print(json.dumps({"ok": False, "error": "没有错题集.xlsx，请提供已有错题集"},
                         ensure_ascii=False))
        return
    ensure_books(root)
    _, ws = open_book(path, W_SHEET)
    _, rows = _wrong_rows(ws)
    board = (args.board or "").strip()
    items, mastered = [], 0
    for it in rows:
        d = it["data"]
        if board and str(d.get("板块") or "").strip() != board:
            continue
        cnt = _as_int(d.get("通过次数"))
        rec = {
            "序号": it["row"] - 1,
            "板块": d.get("板块") or "",
            "题目": d.get("题目") or "",
            "正确答案": d.get("正确答案") or "",
            "通过次数": cnt,
        }
        if cnt >= args.threshold:
            mastered += 1
            if args.all:
                items.append(rec)
        else:
            items.append(rec)
    print(json.dumps({
        "ok": True, "阈值": args.threshold,
        "待复习": sum(1 for i in items if i["通过次数"] < args.threshold),
        "已过关省略": mastered,
        "错题": items,
    }, ensure_ascii=False))


def _update_wrong_count(path: str, board: str, question: str,
                        threshold: int, mode: str) -> dict:
    wb, ws = open_book(path, W_SHEET)
    col, rows = _wrong_rows(ws)
    q = question.strip()
    b = (board or "").strip()
    hit = next((it for it in rows
                if str(it["data"].get("题目") or "").strip() == q
                and (not b or str(it["data"].get("板块") or "").strip() == b)), None)
    if not hit:
        return {"ok": False, "error": f"错题集里没找到题目「{question}」"}
    cur = _as_int(hit["data"].get("通过次数"))
    if mode == "pass":
        new = cur + 1
        action = "复习通过，通过次数 +1"
    else:
        new = 0
        action = "复习答错，通过次数清零"
    ws.cell(row=hit["row"], column=col["通过次数"] + 1, value=new)
    wb.save(path)
    return {"ok": True, "题目": question, "板块": hit["data"].get("板块") or "",
            "通过次数": new, "已过关": new >= threshold, "动作": action}


def cmd_review_pass(args) -> None:
    root = rnsk_root(args.dir)
    ensure_books(root)
    result = _update_wrong_count(os.path.join(root, WRONG_BOOK),
                                 args.board, args.question, args.threshold, "pass")
    print(json.dumps(result, ensure_ascii=False))


def cmd_review_fail(args) -> None:
    root = rnsk_root(args.dir)
    ensure_books(root)
    result = _update_wrong_count(os.path.join(root, WRONG_BOOK),
                                 args.board, args.question, args.threshold, "fail")
    print(json.dumps(result, ensure_ascii=False))


# ---------------------------------------------------------------- 主命令

def cmd_init(args) -> None:
    root = rnsk_root(args.dir)
    idx = os.path.join(root, INDEX_NAME)
    if os.path.exists(idx) and not args.force:
        boards = parse_index(idx)
        print(json.dumps({"ok": True, "already": True, "rnsk": root,
                          "boards": len(boards),
                          "hint": "rnsk 已存在，未覆盖；要重建请加 --force"},
                         ensure_ascii=False))
        return

    os.makedirs(root, exist_ok=True)
    entry = args.url or entry_url(args.lang)
    toc = parse_toc(entry)
    if not toc:
        raise SystemExit(f"没能从 {entry} 解析出目录，请用 --url 指定正确的教程入口")
    if args.limit:
        toc = toc[:args.limit]
    # 首页就是总目录，去掉它自己
    toc = [t for t in toc if t["url"].rstrip("/") != entry.rstrip("/")] or toc

    def work(item: dict) -> dict:
        try:
            info = page_info(item["url"], item["title"])
        except Exception:  # noqa: BLE001
            info = {"heads": [], "desc": ""}
        item["points"] = gen_points(item["title"], info["heads"])
        item["summary"] = info["desc"] or f"{item['title']} 的核心知识点讲解"
        item["category"], item["skip"] = classify_board(item["title"], item["summary"])
        return item

    with ThreadPoolExecutor(max_workers=6) as pool:
        boards = list(pool.map(work, toc))
    filtered = [b for b in boards if b.get("category") != "推广/导读"]

    with open(idx, "w", encoding="utf-8") as fh:
        fh.write(render_index(args.lang or "编程语言", filtered))
    ensure_books(root)

    print(json.dumps({
        "ok": True, "rnsk": root, "entry": entry,
        "boards": len(filtered),
        "filtered_promo": len(boards) - len(filtered),
        "first_board": filtered[0]["title"] if filtered else None,
        "files": [INDEX_NAME, QUESTION_BOOK, WRONG_BOOK],
    }, ensure_ascii=False))


def cmd_next(args) -> None:
    root = rnsk_root(args.dir)
    idx = os.path.join(root, INDEX_NAME)
    if not os.path.exists(idx):
        print(json.dumps({"ok": False, "error": "没有 rnsk/知识目录.md，先执行 init"},
                         ensure_ascii=False))
        return
    boards = parse_index(idx)
    todo = [b for b in boards if not is_done(b)]
    if not todo:
        print(json.dumps({"ok": True, "done": True, "total": len(boards),
                          "message": "所有板块都学完了"}, ensure_ascii=False))
        return
    board = todo[0]
    url = board.get("板块地址", "")
    text, err = "", None
    if url:
        try:
            text = extract_text(url, limit=args.chars)
        except Exception as exc:  # noqa: BLE001
            err = str(exc)
    print(json.dumps({
        "ok": True, "done": False,
        "板块名字": board.get("板块名字", ""),
        "板块地址": url,
        "板块要点": board.get("板块要点", ""),
        "板块内容": board.get("板块内容", ""),
        "剩余未完成": len(todo),
        "总板块数": len(boards),
        "正文": text,
        "抓取错误": err,
    }, ensure_ascii=False))


def cmd_mark(args) -> None:
    root = rnsk_root(args.dir)
    path = os.path.join(root, INDEX_NAME)
    if not os.path.exists(path):
        print(json.dumps({"ok": False, "error": "没有 rnsk/知识目录.md"},
                         ensure_ascii=False))
        return
    target = args.board.strip()
    done_val = "false" if args.undo else "true"
    lines = open(path, encoding="utf-8").read().splitlines()
    out, cur, hit = [], None, False
    for line in lines:
        s = line.strip()
        if s.startswith("[板块名字]:"):
            cur = s.split("]:", 1)[1].strip()
        if cur == target and s.startswith("[是否完成]:"):
            line = f"[是否完成]: {done_val}"
            hit = True
        out.append(line)
    if not hit:
        print(json.dumps({"ok": False, "error": f"没找到板块「{target}」"},
                         ensure_ascii=False))
        return
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(out) + "\n")
    print(json.dumps({"ok": True, "板块": target, "是否完成": done_val},
                     ensure_ascii=False))


def cmd_stats(args) -> None:
    root = rnsk_root(args.dir)
    result: dict = {"ok": True, "rnsk": root, "存在": os.path.isdir(root)}
    idx = os.path.join(root, INDEX_NAME)
    if os.path.exists(idx):
        boards = parse_index(idx)
        done = [b for b in boards if is_done(b)]
        result.update({
            "总板块数": len(boards),
            "已完成": len(done),
            "未完成": len(boards) - len(done),
            "下一个板块": next((b.get("板块名字") for b in boards if not is_done(b)), None),
        })
    for fname, sheet in ((QUESTION_BOOK, Q_SHEET), (WRONG_BOOK, W_SHEET)):
        path = os.path.join(root, fname)
        if os.path.exists(path):
            try:
                _, ws = open_book(path, sheet)
                result[fname] = ws.max_row - 1
                if fname == WRONG_BOOK:
                    _, rows = _wrong_rows(ws)
                    result["待复习错题"] = sum(
                        1 for it in rows
                        if _as_int(it["data"].get("通过次数")) < DEFAULT_REVIEW_THRESHOLD)
            except Exception as exc:  # noqa: BLE001
                result[fname] = f"读取失败: {exc}"
        else:
            result[fname] = None
    print(json.dumps(result, ensure_ascii=False))


# ---------------------------------------------------------------- CLI

def main() -> None:
    ap = argparse.ArgumentParser(description="rnsk-tutor helper")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("init", help="抓菜鸟教程目录，生成知识目录.md 和两个 Excel")
    p.add_argument("--dir", required=True, help="工作目录（rnsk 文件夹会建在这里）")
    p.add_argument("--lang", default="", help="语言名，如 python3 / java / cpp")
    p.add_argument("--url", default="", help="直接指定菜鸟教程入口地址")
    p.add_argument("--limit", type=int, default=0, help="只取前 N 个板块（调试用）")
    p.add_argument("--force", action="store_true", help="已存在时重建")
    p.set_defaults(func=cmd_init)

    p = sub.add_parser("next", help="取下一个未完成板块，并抓回正文")
    p.add_argument("--dir", required=True)
    p.add_argument("--chars", type=int, default=6000, help="正文最多返回多少字符")
    p.set_defaults(func=cmd_next)

    p = sub.add_parser("mark", help="把某板块标记为已完成")
    p.add_argument("--dir", required=True)
    p.add_argument("--board", required=True)
    p.add_argument("--undo", action="store_true", help="改回未完成")
    p.set_defaults(func=cmd_mark)

    p = sub.add_parser("addq", help="往题目册.xlsx 追加一道题")
    p.add_argument("--dir", required=True)
    p.add_argument("--board", required=True)
    p.add_argument("--type", default="单选题", help="单选题/多选题/填空题")
    p.add_argument("--question", required=True)
    p.add_argument("--options", default="", help="选项，用 | 分隔，如 'A.xx|B.xx'")
    p.add_argument("--answer", required=True, help="正确答案，多选写 ABD")
    p.add_argument("--point", default="", help="本题对应知识点")
    p.add_argument("--difficulty", default="基础", help="基础/进阶/综合")
    p.set_defaults(func=cmd_addq)

    p = sub.add_parser("review", help="列出还没过关的错题")
    p.add_argument("--dir", required=True)
    p.add_argument("--board", default="", help="只看某个板块")
    p.add_argument("--threshold", type=int, default=DEFAULT_REVIEW_THRESHOLD,
                   help="连续通过几次算过关（默认 3）")
    p.add_argument("--all", action="store_true", help="连带已过关的错题一起列出")
    p.set_defaults(func=cmd_review)

    for name, fn, helptext in (
        ("review-pass", cmd_review_pass, "复习答对：通过次数 +1"),
        ("review-fail", cmd_review_fail, "复习答错：通过次数清零"),
    ):
        p = sub.add_parser(name, help=helptext)
        p.add_argument("--dir", required=True)
        p.add_argument("--board", default="")
        p.add_argument("--question", required=True)
        p.add_argument("--threshold", type=int, default=DEFAULT_REVIEW_THRESHOLD,
                       help="连续通过几次算过关（默认 3）")
        p.set_defaults(func=fn)

    p = sub.add_parser("stats", help="看学习进度")
    p.add_argument("--dir", required=True)
    p.set_defaults(func=cmd_stats)

    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"},
                         ensure_ascii=False))
        sys.exit(1)
